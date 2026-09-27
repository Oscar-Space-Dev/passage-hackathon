"""Build the two-page Passage hackathon handout for jury review."""
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'Passage_dossier_hackathon.docx'
INK = RGBColor(25, 45, 43)
HEADING_INK = RGBColor(0, 0, 0)


def add_text(document, text, *, bold_lead=None, style=None):
    paragraph = document.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def shade(cell, color='EAF3EF'):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        shade(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = INK
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    styles = document.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    for name, size in [('Title', 21), ('Heading 1', 13), ('Heading 2', 11)]:
        style = styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = HEADING_INK
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)

    document.add_paragraph('Passage pour la recherche doctorale', 'Title')
    add_text(document, 'Dossier de présentation pour le jury du X IA Hackathon Rise of Agents X')
    document.add_heading('Ce que Passage permet', 1)
    add_text(document, 'Passage aide un doctorant à conduire un projet de recherche avec des agents spécialisés. Il relie la lecture des sources, la préparation d’expériences, la rédaction, les simulations et le développement logiciel à des tâches suivies dans un même dossier. L’utilisateur garde la décision scientifique : les agents proposent et exécutent dans le cadre qu’il a approuvé, puis rendent leurs traces et leurs limites visibles.')

    document.add_heading('Un parcours vérifiable', 1)
    add_table(document, ['Étape', 'Ce que fait Passage', 'Décision humaine'], [
        ('Cadrer', 'Créer un projet et choisir une tâche parmi 32 travaux doctoraux R1 à R3.', 'Définir la question, les moyens et les critères de réussite.'),
        ('Mobiliser les agents', 'Marguerite propose un plan et peut créer un spécialiste avec un harnais contrôlé.', 'Approuver le plan avant toute exécution.'),
        ('Produire', 'L’agent prépare un brouillon fondé sur les sources et pièces autorisées.', 'Relire, corriger et valider le livrable.'),
        ('Tracer', 'Conserver versions, preuves, décisions et incidents dans le dossier.', 'Vérifier ce qui a réellement été fait.'),
    ])

    document.add_heading('Démonstration proposée', 1)
    add_text(document, 'Un doctorant ouvre un projet, décrit à Marguerite une revue de littérature, inspecte le plan puis l’approuve. Il ouvre ensuite une tâche, examine les sources et les étapes de validation, et dessine un protocole que Passage formalise en brouillon révisable. La vidéo suit ce parcours avec une voix off Gradium générée pour la candidature.')

    document.add_page_break()
    document.add_heading('Une architecture agentique contrôlée', 1)
    add_text(document, 'Chaque agent associe un cerveau choisi par l’utilisateur et un harnais versionné : mission, skill, contexte, mémoire, outils et règles de validation. ChatGPT via Codex est le choix par défaut. Marguerite suit la doctrine Super Skill Facilitator adaptée à Passage pour créer des agents complets ; les définitions incomplètes sont refusées.')
    add_text(document, 'Gradium assure la transcription et la synthèse vocale. Dust et Pipelex sont reliés par MCP avec des outils explicitement accordés. Le dessin intégré repose sur Excalidraw et peut alimenter un protocole ou un brouillon de méthode MTHDS. Le SDK Jinkō est une piste d’intégration distincte qui doit être validée par un essai réel avant d’être présentée comme opérationnelle.')

    document.add_heading('Ce que le jury peut vérifier', 1)
    add_table(document, ['Action', 'Preuve attendue'], [
        ('Créer un compte personnel', 'Projets et conversations propres au compte.'),
        ('Discuter avec Marguerite', 'Plan affiché avant approbation et résultat de chaque action.'),
        ('Ouvrir un travail doctoral', 'Questions, pièces, versions et contrôles humains visibles.'),
        ('Consulter une notice de thèse', 'Origine et lien vers la source affichés.'),
        ('Dessiner un protocole', 'Schéma enregistré puis brouillon à relire.'),
    ])

    document.add_heading('Limites présentées honnêtement', 1)
    add_text(document, 'Passage ne conduit pas d’expérience physique et ne valide pas un résultat scientifique à la place du chercheur. Une méthode MTHDS exportée n’est pas automatiquement publiée dans Pipelex. Les connexions externes dépendent des comptes et autorisations de chaque utilisateur. Les notices de thèses renvoient vers leurs sources ; les manuscrits ne sont pas réhébergés sans droit.')

    links = [(name, os.environ.get(key, '').strip()) for name, key in (
        ('Application', 'PASSAGE_PUBLIC_URL'), ('Code source', 'PASSAGE_REPOSITORY_URL'),
        ('Vidéo', 'PASSAGE_VIDEO_URL'))]
    links = [(name, url) for name, url in links if url]
    if links:
        document.add_heading('Accès', 1)
        for name, url in links:
            add_text(document, f'{name} : {url}', bold_lead=f'{name} : ')

    document.core_properties.title = 'Passage pour la recherche doctorale'
    document.core_properties.subject = 'Dossier de présentation hackathon'
    document.core_properties.author = 'Équipe Passage'
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
